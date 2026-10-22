import io
import json
import tempfile
import unittest
from pathlib import Path
from openpyxl import Workbook, load_workbook
from app.core import HEADERS, Invalid, parse_workbook, relative_url, parse_json, digest
from app.runner import run_steps
from app.reports import build_reports, safe_cell
from app.store import Store
from tools.make_sample import sample_rows

BASE='http://127.0.0.1:8765'

def book(rows, headers=HEADERS):
    wb=Workbook(); ws=wb.active; ws.append(headers)
    for row in rows: ws.append(row)
    out=io.BytesIO(); wb.save(out); return out.getvalue()

class CoreTests(unittest.TestCase):
    def test_sample_plan(self):
        plan=parse_workbook(book(sample_rows()),BASE)
        self.assertEqual(len(plan['cases']),3)
        self.assertEqual(sum(len(c['steps']) for c in plan['cases']),11)
        self.assertEqual(digest(plan),digest(json.loads(json.dumps(plan))))

    def test_bad_file(self):
        with self.assertRaises(Invalid): parse_workbook(b'not excel', BASE)

    def test_no_assertion(self):
        with self.assertRaises(Invalid):
            parse_workbook(book([['A','M',1,'{"action":"goto","value":"/demo"}',None]]),BASE)

    def test_gap(self):
        rows=sample_rows()[:4]; rows[-1][2]=5
        with self.assertRaises(Invalid): parse_workbook(book(rows),BASE)

    def test_formula(self):
        rows=sample_rows()[:4]; rows[0][3]='=1+1'
        with self.assertRaises(Invalid): parse_workbook(book(rows),BASE)

    def test_numeric_id(self):
        rows=sample_rows()[:4]; rows[0][0]=1
        with self.assertRaises(Invalid): parse_workbook(book(rows),BASE)

    def test_missing_header(self):
        with self.assertRaises(Invalid): parse_workbook(book([],['wrong']),BASE)

    def test_invalid_target(self):
        rows=sample_rows()[:4]
        rows[1][3]='{"action":"fill","target":{"by":"xpath","value":"//input"},"value":"x"}'
        with self.assertRaises(Invalid): parse_workbook(book(rows),BASE)

    def test_bad_variable(self):
        rows=sample_rows()[:4]
        rows[1][3]='{"action":"fill","target":{"by":"label","value":"用户名"},"value":"${AWS_SECRET}"}'
        with self.assertRaises(Invalid): parse_workbook(book(rows),BASE)

    def test_duplicate_json_key(self):
        with self.assertRaises(Invalid): parse_json('{"action":"goto","action":"click"}')

    def test_relative_paths(self):
        self.assertEqual(relative_url(BASE,'/demo'),BASE+'/demo')
        for value in ['https://elsewhere.invalid','//elsewhere.invalid/x','javascript:alert(1)',r'\\evil/x']:
            with self.assertRaises(Invalid): relative_url(BASE,value)

    def test_fail_blocks(self):
        case=parse_workbook(book(sample_rows()),BASE)['cases'][0]
        def execute(page,step,base):
            if step['number']==2: raise AssertionError('secret must not appear')
        result=run_steps(None,case,BASE,lambda p,n:'proof.png',execute)
        self.assertEqual(result['status'],'FAIL')
        self.assertEqual([s['status'] for s in result['steps']],['PASS','FAIL','BLOCKED','BLOCKED'])
        self.assertNotIn('secret',json.dumps(result))

    def test_error_preserved_when_screenshot_fails(self):
        case=parse_workbook(book(sample_rows()),BASE)['cases'][0]
        def execute(*args): raise RuntimeError('sensitive')
        def evidence(*args): raise OSError('disk')
        result=run_steps(None,case,BASE,evidence,execute)
        self.assertEqual(result['status'],'ERROR')
        self.assertEqual(result['steps'][0]['evidence_warning'],'截图采集失败')
        self.assertNotIn('sensitive',json.dumps(result))

    def test_reports(self):
        case=parse_workbook(book(sample_rows()),BASE)['cases'][0]
        case['id']='<script>alert(1)</script>'
        cr=run_steps(None,case,BASE,lambda *a:'',lambda *a:None)
        result={'run_id':'run','plan_hash':'abc','status':'COMPLETED','cases':[cr]}
        with tempfile.TemporaryDirectory() as tmp:
            build_reports(tmp,result)
            page=(Path(tmp)/'report.html').read_text(encoding='utf8')
            self.assertNotIn('<script>',page)
            self.assertIn('&lt;script&gt;',page)
            self.assertEqual(result['summary']['PASS'],1)
            wb=load_workbook(Path(tmp)/'report.xlsx'); self.assertIn('步骤',wb.sheetnames); wb.close()
        self.assertTrue(safe_cell('=1+1').startswith("'"))

    def test_idempotency_and_restart(self):
        plan=parse_workbook(book(sample_rows()),BASE)
        with tempfile.TemporaryDirectory() as tmp:
            store=Store(tmp); p=store.plan(plan)
            rid,first=store.start(p['plan_id'],p['hash'],BASE)
            again,second=store.start(p['plan_id'],p['hash'],BASE)
            self.assertIsNotNone(first); self.assertIsNone(second); self.assertEqual(rid,again)
            with self.assertRaises(Invalid): store.start(p['plan_id'],'wrong',BASE)
            recovered=Store(tmp).get(rid)
            self.assertEqual(recovered['status'],'ERROR')

if __name__=='__main__': unittest.main()
