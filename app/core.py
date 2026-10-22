"""严格读取五列表格，生成唯一且可审计的执行计划。"""
import hashlib
import io
import json
import re
import zipfile
from urllib.parse import urljoin, urlsplit
from openpyxl import load_workbook

HEADERS = ['测试编号', '模块', '步骤', '操作', '期望结果']
MAX_UPLOAD = 2 * 1024 * 1024
MAX_ROWS = 500
VARIABLES = {'TEST_USER', 'TEST_PASSWORD'}
REF = re.compile(r'\$\{([A-Z][A-Z0-9_]*)\}')

class Invalid(ValueError):
    pass

def require(ok, message):
    if not ok:
        raise Invalid(message)

def object_fields(value, allowed, required):
    require(isinstance(value, dict), '必须是 JSON 对象')
    require(set(value) <= set(allowed), '存在未知字段')
    require(set(required) <= set(value), '缺少字段')

def text(value, label):
    require(isinstance(value, str) and 0 < len(value) <= 2000,
            label + ' 必须是非空文本，且不超过 2000 字符')
    return value

def origin(url):
    p = urlsplit(url)
    require(p.scheme in ('http', 'https') and p.hostname and
            not p.username and not p.password, '环境地址必须是 HTTP(S)，不能包含凭据')
    port = p.port or (443 if p.scheme == 'https' else 80)
    return p.scheme, p.hostname.lower(), port

def relative_url(base, value):
    text(value, '路径')
    p = urlsplit(value)
    require(not p.scheme and not p.netloc and '\\' not in value and
            not any(ord(c) < 32 for c in value), '只接受同站相对路径')
    result = urljoin(base.rstrip('/') + '/', value)
    require(origin(result) == origin(base), '路径不属于测试站点')
    return result

def validate_target(target):
    object_fields(target, ('by', 'value', 'role'), ('by', 'value'))
    require(target['by'] in ('label', 'role', 'text', 'testid'), '不支持的定位器')
    text(target['value'], '定位文本')
    if target['by'] == 'role':
        require(target.get('role') in ('button', 'link', 'textbox', 'checkbox',
                'combobox', 'radio', 'heading', 'option', 'tab'), '无效的 role')
    else:
        require('role' not in target, '非 role 定位不接受 role 字段')

def validate_action(action, base):
    object_fields(action, ('action', 'target', 'value'), ('action',))
    name = action['action']
    require(name in ('goto', 'fill', 'click', 'check'), '不支持的动作')
    if name == 'goto':
        require(set(action) == {'action', 'value'}, 'goto 只接受 value')
        relative_url(base, action['value'])
    else:
        validate_target(action.get('target'))
        if name == 'fill':
            require(set(action) == {'action', 'target', 'value'}, 'fill 字段不完整')
            value = text(action['value'], '输入值')
            if '${' in value:
                m = REF.fullmatch(value)
                require(m is not None and m.group(1) in VARIABLES, '不允许的变量引用')
        else:
            require(set(action) == {'action', 'target'}, '动作不接受 value')

def validate_assertion(assertion, base):
    object_fields(assertion, ('type', 'target', 'value'), ('type',))
    name = assertion['type']
    require(name in ('visible', 'text', 'url'), '不支持的断言')
    if name == 'url':
        require(set(assertion) == {'type', 'value'}, 'url 断言字段不完整')
        relative_url(base, assertion['value'])
    else:
        validate_target(assertion.get('target'))
        fields = {'type', 'target'}
        if name == 'text':
            fields.add('value')
            text(assertion.get('value'), '预期文本')
        require(set(assertion) == fields, '断言字段不完整或存在多余字段')

def parse_json(value):
    # 禁止 JSON 重复键、NaN、Infinity，避免不同解析器产生不同计划。
    def pairs(items):
        result = {}
        for k, v in items:
            require(k not in result, 'JSON 存在重复键')
            result[k] = v
        return result
    def constant(_):
        raise Invalid('JSON 不接受 NaN 或 Infinity')
    return json.loads(text(value, 'JSON'), object_pairs_hook=pairs,
                      parse_constant=constant)

def digest(plan):
    raw = json.dumps(plan, ensure_ascii=False, sort_keys=True,
                     separators=(',', ':')).encode('utf-8')
    return hashlib.sha256(raw).hexdigest()

def parse_workbook(data, base):
    origin(base)
    require(0 < len(data) <= MAX_UPLOAD, '文件不能为空，且最大为 2 MiB')
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            info = z.infolist()
            require(len(info) <= 100, '工作簿压缩包文件数超限')
            require(sum(i.file_size for i in info) <= 16 * 1024 * 1024,
                    '工作簿解压后大小超限')
            require('[Content_Types].xml' in z.namelist(), '不是 XLSX 文件')
            require(not any('vbaproject' in i.filename.lower() for i in info),
                    '不接受宏工作簿')
        wb = load_workbook(io.BytesIO(data), read_only=True,
                           data_only=False, keep_links=False)
    except Invalid:
        raise
    except Exception:
        raise Invalid('无法读取 XLSX 文件') from None
    try:
        require(len(wb.worksheets) == 1, '本版本只接受一个工作表')
        ws = wb.worksheets[0]
        # 不信任工作簿声明的维度，按实际 XML 行读取并限流。
        ws.reset_dimensions()
        rows = ws.iter_rows(max_col=6)
        first = next(rows, None)
        require(first is not None, '工作表为空')
        headers = [c.value for c in first]
        require(headers[:5] == HEADERS and headers[5] is None,
                '列名及顺序必须是：' + '、'.join(HEADERS))
        groups = {}
        count = 0
        for row_no, cells in enumerate(rows, 2):
            require(row_no <= MAX_ROWS + 1, '最多 500 个数据行（包含空行）')
            values = [c.value for c in cells]
            if all(v is None for v in values):
                continue
            require(values[5] is None, f'第 {row_no} 行存在额外列')
            require(not any(c.data_type == 'f' for c in cells),
                    f'第 {row_no} 行禁止公式')
            try:
                cid, module, number, raw_action, raw_expected = values[:5]
                require(isinstance(cid, str) and cid.strip() and len(cid) <= 80,
                        '测试编号须按文本保存，最长 80 字符')
                text(module, '模块')
                require(type(number) is int and number > 0, '步骤须为正整数')
                action = parse_json(raw_action)
                validate_action(action, base)
                expected = None
                if raw_expected not in (None, ''):
                    expected = parse_json(raw_expected)
                    validate_assertion(expected, base)
                case = groups.setdefault(cid, {'id': cid, 'module': module, 'steps': []})
                require(case['module'] == module, '同一编号对应了多个模块')
                case['steps'].append({'number': number, 'row': row_no,
                                      'operation': action, 'expected': expected})
                count += 1
            except (ValueError, TypeError, KeyError, AttributeError):
                # 不把原始单元格/输入值回显到错误或日志中。
                raise Invalid(f'第 {row_no} 行无效，请检查动作、定位和断言格式') from None
        require(count > 0, '没有可执行步骤')
        for case in groups.values():
            case['steps'].sort(key=lambda s: s['number'])
            require([s['number'] for s in case['steps']] ==
                    list(range(1, len(case['steps']) + 1)),
                    '存在步骤重复或不连续的用例')
            require(any(s['expected'] for s in case['steps']), '每个用例至少一个断言')
        return {'version': 1, 'base_url': base, 'cases': list(groups.values())}
    finally:
        wb.close()
