import html
import json
from pathlib import Path
from collections import Counter
from openpyxl import Workbook

def safe_cell(value):
    if isinstance(value, str) and value.lstrip().startswith(('=', '+', '-', '@')):
        return "'" + value
    return value

def build_reports(folder, result):
    folder = Path(folder)
    folder.mkdir(parents=True, exist_ok=True)
    counts = Counter(c['status'] for c in result['cases'])
    result['summary'] = {s: counts[s] for s in ('PASS', 'FAIL', 'ERROR', 'BLOCKED')}
    (folder / 'report.json').write_text(json.dumps(result, ensure_ascii=False, indent=2),
                                       encoding='utf-8')
    wb = Workbook()
    summary = wb.active
    summary.title = '摘要'
    summary.append(['任务', result['run_id']])
    summary.append(['状态', result['status']])
    summary.append(['计划哈希', result['plan_hash']])
    for k, v in result['summary'].items():
        summary.append([k, v])
    cases = wb.create_sheet('用例')
    cases.append(['编号', '模块', '状态'])
    steps = wb.create_sheet('步骤')
    steps.append(['编号', '步骤', '源行', '动作', '期望结果', '状态', '实际结果', '耗时秒', '截图', '证据告警'])
    body = []
    for case in result['cases']:
        cases.append([safe_cell(case['id']), safe_cell(case['module']), case['status']])
        for step in case['steps']:
            row = [case['id'], step['number'], step['row'], step.get('action', ''),
                   json.dumps(step.get('expected'), ensure_ascii=False), step['status'],
                   step['message'], step['seconds'], step.get('screenshot', ''),
                   step.get('evidence_warning', '')]
            steps.append([safe_cell(v) for v in row])
            body.append('<tr>' + ''.join('<td>' + html.escape(str(v)) + '</td>'
                                         for v in row) + '</tr>')
    wb.save(folder / 'report.xlsx')
    markup = ('<!doctype html><html lang="zh-CN"><meta charset="utf-8">'
              '<title>自动化测试报告</title><style>body{font-family:sans-serif;padding:24px}'
              'td,th{border:1px solid #ccc;padding:8px}table{border-collapse:collapse}</style>'
              '<h1>自动化测试报告</h1><p>任务：' + html.escape(result['run_id']) +
              '</p><p>状态：' + html.escape(result['status']) + '</p><p>' +
              html.escape(json.dumps(result['summary'])) +
              '</p><p>错误：' + html.escape(result.get('error', '')) +
              '</p><table><tr><th>编号</th><th>步骤</th><th>源行</th><th>动作</th><th>期望结果</th><th>状态</th>'
              '<th>实际结果</th><th>耗时秒</th><th>截图文件</th><th>证据告警</th></tr>' +
              ''.join(body) + '</table></html>')
    (folder / 'report.html').write_text(markup, encoding='utf-8')
